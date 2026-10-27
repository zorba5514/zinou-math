"""ينشئ قالب Excel لاستيراد الأسئلة: public/templates/questions.xlsx
ويُنشئ ملفات تجريبية للاختبارات الآلية في test/fixtures/.

يُشغَّل يدويًا عند تغيير القالب فقط (ليس جزءًا من النشر):
    python3 scripts/make-questions-template.py
الأعمدة يجب أن تطابق src/quiz-import.ts.
"""
import zipfile
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parent.parent
HEADERS = [
    ("النوع", 18),
    ("نص السؤال", 48),
    ("العلامة", 9),
    ("الاختيار 1", 16),
    ("الاختيار 2", 16),
    ("الاختيار 3", 16),
    ("الاختيار 4", 16),
    ("الاختيار 5", 14),
    ("الاختيار 6", 14),
    ("الإجابة الصحيحة", 30),
    ("هامش الخطأ", 12),
    ("الذكاء الاصطناعي", 16),
]
TYPES = ["اختيار واحد", "اختيار متعدد", "صح أو خطأ", "إجابة رقمية", "إجابة نصية", "عبارة رياضية", "سؤال مفتوح"]
EXAMPLES = [
    ["اختيار واحد", "كم يساوي $\\frac{1}{2}+\\frac{1}{4}$ ؟", "2", "$\\frac{2}{6}$", "$\\frac{3}{4}$", "1", "", "", "", "2", "", ""],
    ["اختيار متعدد", "أي هذه الأعداد أولية؟", "2", "2", "4", "7", "9", "", "", "1، 3", "", ""],
    ["صح أو خطأ", "العدد 0 عدد زوجي.", "1", "", "", "", "", "", "", "صح", "", ""],
    ["إجابة رقمية", "احسب $3 \\div 4$", "2", "", "", "", "", "", "", "3/4", "", ""],
    ["إجابة رقمية", "قيمة مقربة لـ $\\sqrt{2}$ إلى 0,01", "1", "", "", "", "", "", "", "1,41", "0,01", ""],
    ["إجابة نصية", "ما اسم الرباعي الذي فيه كل ضلعين متقابلين متوازيين؟", "2", "", "", "", "", "", "", "متوازي الأضلاع\nparallélogramme", "", "نعم"],
    ["عبارة رياضية", "انشر $(x+1)^2$", "3", "", "", "", "", "", "", "x^2+2x+1", "", ""],
    ["سؤال مفتوح", "برهن أن مجموع زوايا المثلث 180°", "5", "", "", "", "", "", "", "", "", ""],
]
GUIDE = [
    ("كيف تملأ القالب", None),
    ("اكتب سؤالًا في كل صف من ورقة «الأسئلة»، ابتداءً من الصف 2. لا تغيّر صف العناوين.", None),
    ("ورقة «أمثلة» فيها سؤال من كل نوع: انسخ منها ما تحتاج.", None),
    ("", None),
    ("النوع", "اختره من القائمة المنسدلة في العمود الأول."),
    ("نص السؤال", "الصيغ الرياضية بين علامتي $ مثل: $\\frac{3}{4}$ أو $x^2$"),
    ("العلامة", "عدد النقاط (فارغ = 1). مثل: 2 أو 1,5"),
    ("الاختيارات", "للاختيار الواحد والمتعدد: اكتب من 2 إلى 6 اختيارات."),
    ("الإجابة الصحيحة", ""),
    ("   اختيار واحد", "رقم الاختيار الصحيح، مثل: 2"),
    ("   اختيار متعدد", "أرقام الاختيارات الصحيحة مفصولة بفاصلة، مثل: 1، 3"),
    ("   صح أو خطأ", "صح أو خطأ"),
    ("   إجابة رقمية", "العدد، مثل: 12 أو 3/4 أو 0,75 أو √2 (يقبل التلميذ أي كتابة مساوية)"),
    ("   إجابة نصية", "كل الإجابات المقبولة، كل واحدة في سطر (Alt+Enter داخل الخانة) أو مفصولة بـ |"),
    ("   عبارة رياضية", "العبارة، مثل: 2x+3 (تُقبل 3+2x و(4x+6)/2 أيضًا)"),
    ("   سؤال مفتوح", "اتركها فارغة: تصححه أنت."),
    ("هامش الخطأ", "للإجابة الرقمية فقط (اختياري)، مثل: 0,01"),
    ("الذكاء الاصطناعي", "للإجابة النصية: «لا» لإطفائه. فارغ أو «نعم» = يحكم على الإجابات غير الموجودة في القائمة."),
    ("", None),
    ("بعد الرفع تظهر لك معاينة الأسئلة والأخطاء إن وُجدت، ولا يُضاف شيء قبل أن تؤكد.", None),
]

INK = "16181D"
BRAND = "FBBF24"
thin = Side(style="thin", color="D1D5DB")


def question_sheet(ws, rows):
    ws.sheet_view.rightToLeft = True
    ws.freeze_panes = "A2"
    for i, (title, width) in enumerate(HEADERS, start=1):
        cell = ws.cell(row=1, column=i, value=title)
        cell.font = Font(bold=True, color=INK, size=12)
        cell.fill = PatternFill("solid", fgColor=BRAND)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(bottom=Side(style="medium", color=INK))
        ws.column_dimensions[cell.column_letter].width = width
    ws.row_dimensions[1].height = 30
    # كل الخانات نصية، حتى لا يحوّل Excel «3/4» إلى تاريخ أو «0,75» إلى رقم بصيغة أخرى
    for r in range(2, 202):
        for c in range(1, len(HEADERS) + 1):
            cell = ws.cell(row=r, column=c)
            cell.number_format = "@"
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = Border(bottom=thin)
    for r, values in enumerate(rows, start=2):
        for c, v in enumerate(values, start=1):
            ws.cell(row=r, column=c, value=v or None)
    types = DataValidation(type="list", formula1='"' + ",".join(TYPES) + '"', allow_blank=True)
    types.error = "اختر النوع من القائمة."
    types.errorTitle = "نوع غير معروف"
    types.prompt = "اختر نوع السؤال"
    ws.add_data_validation(types)
    types.add("A2:A201")
    ai = DataValidation(type="list", formula1='"نعم,لا"', allow_blank=True)
    ws.add_data_validation(ai)
    ai.add("L2:L201")


def build(path, main_rows):
    wb = Workbook()
    ws = wb.active
    ws.title = "الأسئلة"
    question_sheet(ws, main_rows)
    question_sheet(wb.create_sheet("أمثلة"), EXAMPLES)
    guide = wb.create_sheet("طريقة الملء")
    guide.sheet_view.rightToLeft = True
    guide.column_dimensions["A"].width = 22
    guide.column_dimensions["B"].width = 90
    for r, (a, b) in enumerate(GUIDE, start=1):
        ca = guide.cell(row=r, column=1, value=a or None)
        ca.alignment = Alignment(wrap_text=b is not None, vertical="top")
        if b is None and a:
            guide.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)
        if r == 1:
            ca.font = Font(bold=True, size=14)
        elif b is not None:
            ca.font = Font(bold=not a.startswith("   "))
            guide.cell(row=r, column=2, value=b or None).alignment = Alignment(wrap_text=True, vertical="top")
    wb.properties.creator = "منصة الأستاذ زين الدين للرياضيات"
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def inline_strings_fixture(path):
    """ملف مكتوب يدويًا بخلايا inlineStr وبدون sharedStrings، كما تكتبه بعض البرامج."""
    def cell(ref, text):
        return f'<c r="{ref}" t="inlineStr"><is><t xml:space="preserve">{text}</t></is></c>'
    rows = [
        "".join(cell(f"{c}1", h) for c, h in zip("ABCDEFGHIJKL", [h for h, _ in HEADERS])),
        cell("A2", "صح أو خطأ") + cell("B2", "المربع معين &amp; مستطيل") + '<c r="C2"><v>1.5</v></c>' + cell("J2", "صحيح"),
        cell("A3", "إجابة نصية") + cell("B3", "اسم <r><t>الشكل</t></r>") + cell("J3", "دائرة | cercle"),
    ]
    sheet = (
        '<?xml version="1.0" encoding="UTF-8"?><worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><sheetData>'
        + "".join(f'<row r="{i}">{r}</row>' for i, r in enumerate(rows, start=1))
        + "</sheetData></worksheet>"
    ).replace("اسم <r><t>الشكل</t></r>", "ما اسم هذا الشكل؟")
    files = {
        "[Content_Types].xml": '<?xml version="1.0"?><Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"/>',
        "xl/workbook.xml": '<?xml version="1.0"?><workbook xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"><sheets><sheet name="Sheet1" sheetId="1" r:id="rId1"/></sheets></workbook>',
        "xl/_rels/workbook.xml.rels": '<?xml version="1.0"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="worksheet" Target="/xl/worksheets/sheet1.xml"/></Relationships>',
        "xl/worksheets/sheet1.xml": sheet,
    }
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        for name, data in files.items():
            z.writestr(name, data)


def bomb_fixture(path):
    """قنبلة zip صغيرة: ورقة تنتفخ إلى 30 ميغابايت بعد فك الضغط."""
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("xl/workbook.xml", '<workbook><sheets><sheet name="a" r:id="rId1"/></sheets></workbook>')
        z.writestr("xl/_rels/workbook.xml.rels", '<Relationships><Relationship Id="rId1" Target="worksheets/sheet1.xml"/></Relationships>')
        z.writestr("xl/worksheets/sheet1.xml", "<worksheet>" + " " * (30 * 1024 * 1024) + "</worksheet>")


if __name__ == "__main__":
    build(ROOT / "public/templates/questions.xlsx", [])
    fixtures = ROOT / "test/fixtures"
    build(fixtures / "questions-filled.xlsx", EXAMPLES + [
        ["اختيار واحد", "سؤال بلا إجابة صحيحة", "1", "أ", "ب", "", "", "", "", "", "", ""],
        ["", "", "", "", "", "", "", "", "", "", "", ""],
        ["نوع غريب", "سؤال", "1", "", "", "", "", "", "", "", "", ""],
        ["إجابة رقمية", "احسب", "1", "", "", "", "", "", "", "ثلاثة", "", ""],
    ])
    inline_strings_fixture(fixtures / "questions-inline.xlsx")
    bomb_fixture(fixtures / "zip-bomb.xlsx")
    print("ok")
